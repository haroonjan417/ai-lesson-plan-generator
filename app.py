import json
from io import BytesIO

import streamlit as st
from docx import Document

from config import GROQ_API_KEY
from prompts.lesson_prompt import build_lesson_prompt
from services.groq_service import GroqService
def create_word_document(lesson_plan, av_aids=None):
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from io import BytesIO

    document = Document()

    # ---------------------------------------------------------
    # PAGE SETUP - A4 PORTRAIT
    # ---------------------------------------------------------
    section = document.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)

    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.50)
    section.right_margin = Inches(0.50)

    # ---------------------------------------------------------
    # GLOBAL FONT
    # ---------------------------------------------------------
    styles = document.styles

    normal_style = styles["Normal"]
    normal_style.font.name = "Arial"
    normal_style.font.size = Pt(10)

    # ---------------------------------------------------------
    # HELPER FUNCTIONS
    # ---------------------------------------------------------
    def set_cell_shading(cell, fill):
        tc_pr = cell._tc.get_or_add_tcPr()

        shd = tc_pr.find(qn("w:shd"))

        if shd is None:
            shd = OxmlElement("w:shd")
            tc_pr.append(shd)

        shd.set(qn("w:fill"), fill)

    def set_cell_margins(cell, top=70, start=80, bottom=70, end=80):
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()

        tcMar = tcPr.first_child_found_in("w:tcMar")

        if tcMar is None:
            tcMar = OxmlElement("w:tcMar")
            tcPr.append(tcMar)

        for margin, value in [
            ("top", top),
            ("start", start),
            ("bottom", bottom),
            ("end", end),
        ]:
            node = tcMar.find(qn(f"w:{margin}"))

            if node is None:
                node = OxmlElement(f"w:{margin}")
                tcMar.append(node)

            node.set(qn("w:w"), str(value))
            node.set(qn("w:type"), "dxa")

    def set_cell_text(cell, text, bold=False, size=10):
        cell.text = ""

        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0

        run = paragraph.add_run(str(text))
        run.bold = bold
        run.font.name = "Arial"
        run.font.size = Pt(size)

        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

        set_cell_margins(cell)

    def set_repeat_table_header(row):
        trPr = row._tr.get_or_add_trPr()
        tblHeader = OxmlElement("w:tblHeader")
        tblHeader.set(qn("w:val"), "true")
        trPr.append(tblHeader)

    def set_table_borders(table):
        tbl = table._tbl
        tblPr = tbl.tblPr

        borders = tblPr.first_child_found_in("w:tblBorders")

        if borders is None:
            borders = OxmlElement("w:tblBorders")
            tblPr.append(borders)

        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            tag = f"w:{edge}"
            element = borders.find(qn(tag))

            if element is None:
                element = OxmlElement(tag)
                borders.append(element)

            element.set(qn("w:val"), "single")
            element.set(qn("w:sz"), "6")
            element.set(qn("w:space"), "0")
            element.set(qn("w:color"), "B7B7B7")

    def add_section_heading(text):
        paragraph = document.add_paragraph()

        paragraph.paragraph_format.space_before = Pt(5)
        paragraph.paragraph_format.space_after = Pt(2)

        run = paragraph.add_run(text.upper())
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(11)

        return paragraph

    def add_compact_text(text, size=9.5):
        paragraph = document.add_paragraph()

        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1.0

        run = paragraph.add_run(str(text))
        run.font.name = "Arial"
        run.font.size = Pt(size)

        return paragraph

    # ---------------------------------------------------------
    # GET LESSON INFORMATION
    # ---------------------------------------------------------
    lesson_info = lesson_plan.get("lesson_information", {})

    school_name = lesson_info.get("school_name", "")
    teacher_name = lesson_info.get("teacher_name", "")
    lesson_date = lesson_info.get("lesson_date", "")
    curriculum = lesson_info.get("curriculum", "")
    grade = lesson_info.get("grade", "")
    subject = lesson_info.get("subject", "")
    topic = lesson_info.get("topic", "")
    duration = lesson_info.get("duration_minutes", "")
    language = lesson_info.get("language", "")

    # ---------------------------------------------------------
    # TITLE
    # ---------------------------------------------------------
    title = document.add_paragraph()

    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(6)

    run = title.add_run("LESSON PLAN")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(16)

    # ---------------------------------------------------------
    # LESSON INFORMATION TABLE
    # ---------------------------------------------------------
    info_table = document.add_table(rows=5, cols=4)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_table.autofit = False

    info_widths = [
        Inches(0.85),
        Inches(3.05),
        Inches(0.85),
        Inches(3.05),
    ]

    info_data = [
        ("School", school_name, "Date", lesson_date),
        ("Teacher", teacher_name, "Class", grade),
        ("Subject", subject, "Topic", topic),
        ("Curriculum", curriculum, "Duration", f"{duration} min"),
        ("Language", language, "AV / Teaching Aids",
         ", ".join(av_aids) if av_aids else "None specified"),
    ]

    for row_index, row_data in enumerate(info_data):

        row = info_table.rows[row_index]

        for col_index, value in enumerate(row_data):

            cell = row.cells[col_index]
            cell.width = info_widths[col_index]

            if col_index in [0, 2]:
                set_cell_shading(cell, "EDEDED")
                set_cell_text(
                    cell,
                    value,
                    bold=True,
                    size=9
                )
            else:
                set_cell_text(
                    cell,
                    value,
                    bold=False,
                    size=9
                )

    set_table_borders(info_table)

    # ---------------------------------------------------------
    # LEARNING OBJECTIVES
    # ---------------------------------------------------------
    add_section_heading("Learning Objectives")

    objectives = lesson_plan.get("learning_objectives", [])

    if isinstance(objectives, list):

        for i, objective in enumerate(objectives, start=1):
            paragraph = document.add_paragraph()

            paragraph.paragraph_format.left_indent = Inches(0.10)
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(1)
            paragraph.paragraph_format.line_spacing = 1.0

            run = paragraph.add_run(
                f"{i}. {objective}"
            )

            run.font.name = "Arial"
            run.font.size = Pt(9.5)

    else:
        add_compact_text(objectives)

    # ---------------------------------------------------------
    # PRIOR KNOWLEDGE
    # ---------------------------------------------------------
    add_section_heading("Prior Knowledge")

    prior_knowledge = lesson_plan.get(
        "prior_knowledge",
        ""
    )

    add_compact_text(
        prior_knowledge,
        size=9.5
    )

    # ---------------------------------------------------------
    # LESSON SEQUENCE
    # ---------------------------------------------------------
    add_section_heading("Lesson Sequence")

    sequence = lesson_plan.get(
        "lesson_sequence",
        []
    )

    sequence_table = document.add_table(
        rows=1,
        cols=5
    )

    sequence_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sequence_table.autofit = False

    headers = [
        "Time",
        "Stage",
        "Teacher / Learning Activity",
        "Student Activity",
        "Assessment"
    ]

    column_widths = [
        Inches(0.60),
        Inches(1.00),
        Inches(2.55),
        Inches(1.85),
        Inches(1.22),
    ]

    header_row = sequence_table.rows[0]

    for i, header in enumerate(headers):

        cell = header_row.cells[i]
        cell.width = column_widths[i]

        set_cell_shading(cell, "D9E2F3")

        set_cell_text(
            cell,
            header,
            bold=True,
            size=8.5
        )

        header.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    set_repeat_table_header(header_row)

    for stage in sequence:

        row = sequence_table.add_row()

        values = [
            f"{stage.get('duration_minutes', '')} min",
            stage.get("stage", ""),
            stage.get(
                "teacher_activity",
                stage.get(
                    "teacher_learning_activity",
                    ""
                )
            ),
            stage.get(
                "student_activity",
                ""
            ),
            stage.get(
                "assessment",
                ""
            ),
        ]

        for i, value in enumerate(values):

            cell = row.cells[i]
            cell.width = column_widths[i]

            set_cell_text(
                cell,
                value,
                bold=False,
                size=8.5
            )

            if i == 0:
                cell.paragraphs[0].alignment = (
                    WD_ALIGN_PARAGRAPH.CENTER
                )

    set_table_borders(sequence_table)

    # ---------------------------------------------------------
    # ASSESSMENT / DIFFERENTIATION / HOMEWORK
    # ---------------------------------------------------------
    add_section_heading(
        "Assessment | Differentiation | Homework"
    )

    bottom_table = document.add_table(
        rows=1,
        cols=3
    )

    bottom_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    bottom_table.autofit = False

    bottom_headers = [
        "Assessment",
        "Differentiation",
        "Homework"
    ]

    bottom_widths = [
        Inches(2.42),
        Inches(2.42),
        Inches(2.38)
    ]

    bottom_values = [
        lesson_plan.get("assessment", ""),
        lesson_plan.get("differentiation", ""),
        lesson_plan.get("homework", "")
    ]

    for i in range(3):

        cell = bottom_table.rows[0].cells[i]

        cell.width = bottom_widths[i]

        set_cell_shading(
            cell,
            "EDEDED"
        )

        set_cell_text(
            cell,
            bottom_headers[i],
            bold=True,
            size=9
        )

        cell.paragraphs[0].alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

    content_row = bottom_table.add_row()

    for i in range(3):

        cell = content_row.cells[i]

        cell.width = bottom_widths[i]

        set_cell_text(
            cell,
            bottom_values[i],
            bold=False,
            size=8.5
        )

    set_table_borders(bottom_table)

    # ---------------------------------------------------------
    # TEACHER NOTES
    # ---------------------------------------------------------
    teacher_notes = lesson_plan.get(
        "teacher_notes",
        ""
    )

    if teacher_notes:

        add_section_heading("Teacher Notes")

        add_compact_text(
            teacher_notes,
            size=9
        )

    # ---------------------------------------------------------
    # DURATION VALIDATION
    # ---------------------------------------------------------
    total_duration = sum(
        stage.get("duration_minutes", 0)
        for stage in sequence
    )

    validation_paragraph = document.add_paragraph()

    validation_paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.RIGHT
    )

    validation_paragraph.paragraph_format.space_before = Pt(3)
    validation_paragraph.paragraph_format.space_after = Pt(0)

    if total_duration == duration:

        validation_text = (
            f"Duration validated: "
            f"{total_duration} minutes"
        )

    else:

        validation_text = (
            f"Duration: {total_duration} minutes "
            f"(selected: {duration} minutes)"
        )

    validation_run = validation_paragraph.add_run(
        validation_text
    )

    validation_run.bold = True
    validation_run.font.name = "Arial"
    validation_run.font.size = Pt(8.5)

    # ---------------------------------------------------------
    # SAVE TO MEMORY STREAM
    # ---------------------------------------------------------
    file_stream = BytesIO()

    document.save(file_stream)

    file_stream.seek(0)

    return file_stream

def format_av_aids(av_aids):
    """Format selected AV aids for display."""

    if not av_aids:
        return "None specified"

    return ", ".join(av_aids)
# ---------------------------------------------------------
# PAGE CONFIGURATION
# ---------------------------------------------------------

st.set_page_config(
    page_title="AI Lesson Plan Generator",
    page_icon="📚",
    layout="wide"
)


# ---------------------------------------------------------
# HEADER
# ---------------------------------------------------------

st.title("📚 AI Lesson Plan Generator")

st.write(
    "Create professional, teacher-ready lesson plans using AI."
)

st.divider()


# ---------------------------------------------------------
# API CONFIGURATION CHECK
# ---------------------------------------------------------

if not GROQ_API_KEY:
    st.error(
        "⚠️ Groq API key is not configured. "
        "Please configure GROQ_API_KEY before generating "
        "a lesson plan."
    )
    st.stop()


# ---------------------------------------------------------
# LESSON INPUT FORM
# ---------------------------------------------------------

st.subheader("🏫 School & Teacher Information")

school_name = st.text_input(
    "School Name",
    placeholder="e.g. Government High School"
)

teacher_name = st.text_input(
    "Teacher Name",
    placeholder="e.g. Muhammad Haroon Jan"
)

lesson_date = st.date_input(
    "Lesson Date"
)


st.subheader("📝 Lesson Information")

col1, col2 = st.columns(2)
with col1:

    curriculum = st.text_input(
        "Curriculum",
        placeholder="e.g. Punjab Curriculum / Federal Board / Cambridge"
    )

    grade = st.text_input(
        "Grade / Class",
        placeholder="e.g. Grade 5"
    )

    subject = st.text_input(
        "Subject",
        placeholder="e.g. Mathematics"
    )


with col2:

    topic = st.text_input(
        "Lesson Topic",
        placeholder="e.g. Fractions"
    )

    duration = st.number_input(
        "Class Duration (minutes)",
        min_value=10,
        max_value=180,
        value=40,
        step=5
    )

    language = st.selectbox(
        "Lesson Language",
        [
            "English",
            "Urdu",
            "English + Urdu"
        ]
    )

    av_aids = st.multiselect(
        "🎬 AV / Teaching Aids",
        [
            "Slides / PowerPoint",
            "Video",
            "Projector",
            "Whiteboard",
            "Marker",
            "Blackboard",
            "Textbook",
            "Pictures / Charts",
            "Computer / Laptop",
            "Audio",
            "Other"
        ],
        placeholder="Select the teaching aids you will use"
    )

st.divider()


# ---------------------------------------------------------
# GENERATE BUTTON
# ---------------------------------------------------------

generate_button = st.button(
    "✨ Generate Lesson Plan",
    type="primary",
    use_container_width=True
)


# ---------------------------------------------------------
# GENERATION PROCESS
# ---------------------------------------------------------

if generate_button:

    # ---------------------------------------------
    # VALIDATE REQUIRED FIELDS
    # ---------------------------------------------

    if not curriculum.strip():
        st.warning("Please enter the curriculum.")

    elif not grade.strip():
        st.warning("Please enter the grade/class.")

    elif not subject.strip():
        st.warning("Please enter the subject.")

    elif not topic.strip():
        st.warning("Please enter the lesson topic.")

    else:

        try:

            # -----------------------------------------
            # BUILD PROMPT
            # -----------------------------------------

            prompt = build_lesson_prompt(
    curriculum=curriculum,
    grade=grade,
    subject=subject,
    duration=duration,
    topic=topic,
    language=language,
    school_name=school_name,
    teacher_name=teacher_name,
    lesson_date=str(lesson_date),
    av_aids=av_aids
)

            # -----------------------------------------
            # CONNECT TO GROQ
            # -----------------------------------------

            groq_service = GroqService(
                api_key=GROQ_API_KEY
            )


            # -----------------------------------------
            # GENERATE LESSON PLAN
            # -----------------------------------------

            with st.spinner(
                "🤖 Generating your lesson plan..."
            ):

                response = groq_service.generate_response(
                    prompt=prompt
                )


            # -----------------------------------------
            # PARSE JSON RESPONSE
            # -----------------------------------------

            try:

                lesson_plan = json.loads(response)

            except json.JSONDecodeError:

                st.error(
                    "❌ The AI returned an invalid lesson "
                    "plan format."
                )

                st.code(response)

                st.stop()


            # -----------------------------------------
            # SUCCESS MESSAGE
            # -----------------------------------------

            st.success(
                "✅ Lesson plan generated successfully!"
            )


            # -----------------------------------------
            # LESSON INFORMATION
            # -----------------------------------------

            st.subheader("📚 Lesson Plan")

            lesson_info = lesson_plan["lesson_information"]
          
            # -----------------------------------------
            # ONE-PAGE LESSON PLAN
            # -----------------------------------------

            st.markdown(
                "<h1 style='text-align: center;'>LESSON PLAN</h1>",
                unsafe_allow_html=True
            )

            # -----------------------------------------
            # LESSON HEADER
            # -----------------------------------------

            header_col1, header_col2 = st.columns(2)

            with header_col1:

                st.markdown(
                    f"**🏫 School:** "
                    f"{lesson_info['school_name']}"
                )

                st.markdown(
                    f"**👨‍🏫 Teacher:** "
                    f"{lesson_info['teacher_name']}"
                )

                st.markdown(
                    f"**📚 Curriculum:** "
                    f"{lesson_info['curriculum']}"
                )

                st.markdown(
                    f"**📖 Subject:** "
                    f"{lesson_info['subject']}"
                )

                st.markdown(
                    f"**🎓 Class:** "
                    f"{lesson_info['grade']}"
                )

            with header_col2:

                st.markdown(
                    f"**📅 Date:** "
                    f"{lesson_info['lesson_date']}"
                )

                st.markdown(
                    f"**📝 Topic:** "
                    f"{lesson_info['topic']}"
                )

                st.markdown(
                    f"**⏱️ Duration:** "
                    f"{lesson_info['duration_minutes']} minutes"
                )

                st.markdown(
                    f"**🌐 Language:** "
                    f"{lesson_info['language']}"
                )

                st.markdown(
                    f"**🎬 AV / Teaching Aids:** "
                    f"{format_av_aids(av_aids)}"
                )

            st.divider()

            # -----------------------------------------
            # LEARNING OBJECTIVES
            # -----------------------------------------

            st.subheader("🎯 Learning Objectives")

            objectives_text = " • ".join(
                lesson_plan["learning_objectives"]
            )

            st.write(objectives_text)

            # -----------------------------------------
            # PRIOR KNOWLEDGE
            # -----------------------------------------

            st.subheader("🧠 Prior Knowledge")

            st.write(
                lesson_plan["prior_knowledge"]
            )

            # -----------------------------------------
            # LESSON SEQUENCE
            # -----------------------------------------

            st.subheader("📋 Lesson Sequence")

            sequence_data = []

            for stage in lesson_plan["lesson_sequence"]:

                sequence_data.append(
                    {
                        "Time": (
                            f"{stage['duration_minutes']} min"
                        ),
                        "Stage": stage["stage"],
                        "Teacher / Learning Activity": (
                            stage["teacher_activity"]
                        ),
                        "Student Activity": (
                            stage["student_activity"]
                        ),
                        "Assessment": (
                            stage["assessment_check"]
                        )
                    }
                )

            st.table(sequence_data)

            # -----------------------------------------
            # ASSESSMENT & DIFFERENTIATION
            # -----------------------------------------

            assessment_col, differentiation_col = st.columns(2)

            with assessment_col:

                st.subheader("📝 Assessment")

                formative = lesson_plan[
                    "assessment"
                ]["formative"]

                for item in formative:
                    st.markdown(f"- {item}")

                st.markdown(
                    "**Summative:** "
                    + lesson_plan["assessment"]["summative"]
                )

            with differentiation_col:

                st.subheader("🔄 Differentiation")

                st.markdown(
                    "**Support:** "
                    + lesson_plan["differentiation"]["support"]
                )

                st.markdown(
                    "**Extension:** "
                    + lesson_plan["differentiation"]["extension"]
                )

            # -----------------------------------------
            # HOMEWORK & TEACHER NOTES
            # -----------------------------------------

            homework_col, notes_col = st.columns(2)

            with homework_col:

                st.subheader("🏠 Homework")

                st.write(
                    lesson_plan["homework"]
                )

            with notes_col:

                st.subheader("👨‍🏫 Teacher Notes")

                st.write(
                    lesson_plan["teacher_notes"]
                )

            # -----------------------------------------
            # DURATION VALIDATION
            # -----------------------------------------

            total_stage_duration = sum(
                stage["duration_minutes"]
                for stage in lesson_plan["lesson_sequence"]
            )

            if total_stage_duration == duration:

                st.success(
                    f"⏱️ Duration validated: "
                    f"{total_stage_duration} minutes "
                    f"(matches the selected duration)."
                )

            else:

                st.warning(
                    f"⚠️ Duration mismatch: "
                    f"Selected duration = {duration} minutes, "
                    f"but lesson stages total "
                    f"{total_stage_duration} minutes."
                )            # -----------------------------------------
            # WORD DOWNLOAD
            # -----------------------------------------

            st.divider()

            st.subheader("📥 Download Lesson Plan")

            word_file = create_word_document(
                lesson_plan,
                av_aids=av_aids
            )

            st.download_button(
                label="📄 Download as Word",
                data=word_file,
                file_name=(
                    f"{lesson_info['topic']}_Lesson_Plan.docx"
                ),
                mime=(
                    "application/vnd.openxmlformats-officedocument."
                    "wordprocessingml.document"
                ),
                use_container_width=True
            )


        except Exception as e:

            st.error(
                "❌ An error occurred while generating "
                "the lesson plan."
            )

            st.exception(e)
